"""
Script de index documents tu folder knowledge_base vao Pinecone
Emotion chatbot - Psychology knowledge base
Chay: python scripts/index_documents.py
"""

import os
import sys
from pathlib import Path
from typing import List, Dict
import asyncio
import re
import time
import unicodedata
import hashlib

# Add backend to path
sys.path.insert(0, str(Path(__file__).parent.parent))

try:
    from PyPDF2 import PdfReader
except ImportError:
    print("Lỗi: Chưa cài PyPDF2. Chạy: pip install PyPDF2")
    sys.exit(1)

try:
    from docx import Document as DocxDocument
except ImportError:
    print("Lỗi: Chưa cài python-docx. Chạy: pip install python-docx")
    sys.exit(1)

from backend.ai_service.chatbot.embedding import get_embedding_manager

def remove_accents(text: str) -> str:
    """Remove Vietnamese accents for printing"""
    return ''.join(c for c in unicodedata.normalize('NFD', text) if unicodedata.category(c) != 'Mn')


class DocumentIndexer:
    """Index PDF va Word documents vao Pinecone cho emotion chatbot"""
    
    def __init__(self, knowledge_base_path: str = "knowledge_base"):
        self.knowledge_base_path = knowledge_base_path
        self.chunk_size = 400  # Optimized for Vietnamese text
        self.overlap = 100      # Reduced overlap
        
    def get_all_documents(self) -> List[Dict]:
        """Lấy tất cả PDF và DOCX files"""
        documents = []
        
        for root, dirs, files in os.walk(self.knowledge_base_path):
            for file in files:
                if file.endswith(('.pdf', '.docx')):
                    file_path = os.path.join(root, file)
                    category = os.path.basename(root)  # psychology, mindfulness, emotions
                    
                    documents.append({
                        'path': file_path,
                        'filename': file,
                        'category': category,
                        'type': 'pdf' if file.endswith('.pdf') else 'docx'
                    })
        
        return documents
    
    def parse_pdf(self, file_path: str) -> str:
        """Parse PDF file"""
        try:
            with open(file_path, 'rb') as file:
                reader = PdfReader(file)
                text = ""
                
                if len(reader.pages) == 0:
                    print(f"  Canh bao: PDF rong hoac khong co trang")
                    return ""
                
                for page_num, page in enumerate(reader.pages):
                    try:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n"
                    except Exception as page_error:
                        print(f"  Loi doc trang {page_num + 1}: {page_error}")
                        continue
                
                if not text.strip():
                    print(f"  Canh bao: Khong extract duoc text tu PDF")
                    return ""
                
                return self.clean_text(text)
                
        except FileNotFoundError:
            print(f"  Lỗi: Không tìm thấy file {file_path}")
            return ""
        except Exception as e:
            print(f"  Loi parse PDF {os.path.basename(file_path)}: {str(e)}")
            return ""
    
    def parse_docx(self, file_path: str) -> str:
        """Parse Word document"""
        try:
            doc = DocxDocument(file_path)
            text = ""
            
            # Paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"
            
            # Tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text += " | ".join(row_text) + "\n"
            
            if not text.strip():
                print(f"  Canh bao: Document rong hoac khong co noi dung")
                return ""
            
            return self.clean_text(text)
            
        except FileNotFoundError:
            print(f"  Lỗi: Không tìm thấy file {file_path}")
            return ""
        except Exception as e:
            print(f"  Loi parse DOCX {os.path.basename(file_path)}: {str(e)}")
            return ""
    
    def clean_text(self, text: str) -> str:
        """Clean extracted text"""
        if not text:
            return ""
        
        # Remove extra whitespace but preserve line breaks for structure
        text = re.sub(r'[ \t]+', ' ', text)
        text = re.sub(r'\n\s*\n', '\n\n', text)
        
        # Remove control characters
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', text)
        
        return text.strip()
    
    def chunk_text(self, text: str) -> List[str]:
        """Split text into overlapping chunks"""
        if not text or len(text.strip()) < 50:
            return []
        
        words = text.split()
        
        if len(words) < 50:
            return [text]
        
        chunks = []
        
        for i in range(0, len(words), self.chunk_size - self.overlap):
            chunk_words = words[i:i + self.chunk_size]
            chunk_text = ' '.join(chunk_words)
            
            # Only keep chunks with meaningful content
            if len(chunk_text.strip()) > 100:
                chunks.append(chunk_text.strip())
        
        return chunks
    
    async def index_documents(self):
        """Main indexing function"""
        print("=" * 60)
        print("DOCUMENT INDEXER - Emotion Psychology Knowledge Base")
        print("=" * 60)
        
        # Get all documents
        documents = self.get_all_documents()
        
        if not documents:
            print("\nKhong tim thay documents nao trong knowledge_base/")
            print("Hay them PDF hoac DOCX files vao:")
            print("  - knowledge_base/psychology/")
            print("  - knowledge_base/mindfulness/")
            print("  - knowledge_base/emotions/")
            return
        
        print(f"\nTim thay {len(documents)} documents:")
        for doc in documents:
            print(f"  - {remove_accents(doc['filename'])} ({doc['category']})")
        
        # Initialize embedding manager
        print("\nKhoi tao Pinecone...")
        try:
            embedding_manager = await get_embedding_manager()
        except Exception as e:
            print(f"Loi ket noi Pinecone: {str(e)}")
            print("Kiem tra:")
            print("  1. PINECONE_API_KEY trong .env")
            print("  2. PINECONE_INDEX_NAME trong .env")
            print("  3. Ket noi internet")
            return
        
        # Process each document
        total_chunks = 0
        
        for idx, doc in enumerate(documents, 1):
            print(f"\n[{idx}/{len(documents)}] Processing: {remove_accents(doc['filename'])}")
            
            # Parse document
            if doc['type'] == 'pdf':
                text = self.parse_pdf(doc['path'])
            else:
                text = self.parse_docx(doc['path'])
            
            if not text:
                print(f"  Bo qua: Khong the parse document")
                continue
            
            print(f"  Extracted {len(text)} characters")
            
            # Chunk text
            chunks = self.chunk_text(text)
            
            if not chunks:
                print(f"  bo qua: text qua ngan de chunk")
                continue
            
            print(f"  Created {len(chunks)} chunks")
            
            # Prepare for embedding
            documents_to_embed = []
            for chunk_idx, chunk in enumerate(chunks):
                # Generate stable ID using hash (prevent duplicates)
                stable_id = hashlib.md5(f"{doc['filename']}_{chunk_idx}_{chunk[:100]}".encode()).hexdigest()[:16]
                
                documents_to_embed.append({
                    'id': f"{doc['filename']}_{chunk_idx}_{stable_id}",
                    'content': chunk,
                    'metadata': {
                        'filename': doc['filename'],
                        'category': doc['category'],
                        'chunk_index': chunk_idx,
                        'total_chunks': len(chunks),
                        'file_type': doc['type']
                    }
                })
            
            # Embed and store
            try:
                success = await embedding_manager.embed_texts(documents_to_embed)
                
                if success:
                    print(f"  Indexed {len(chunks)} chunks vào Pinecone")
                    total_chunks += len(chunks)
                else:
                    print(f"  Loi: Khong the index")
            except Exception as e:
                print(f"  Loi khi index: {str(e)}")
        
        print("\n" + "=" * 60)
        print(f"HOAN THANH!")
        print(f"Tong: {len(documents)} documents, {total_chunks} chunks da index")
        print("=" * 60)


async def main():
    """Run indexer"""
    indexer = DocumentIndexer()
    await indexer.index_documents()


if __name__ == "__main__":
    asyncio.run(main())

